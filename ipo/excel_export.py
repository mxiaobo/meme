from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


def build_settlement_workbook(month, lines):
    wb = Workbook()
    ws = wb.active
    ws.title = f"结算 {month}"

    headers = [
        "姓名", "本月净利润", "应得分红(30%)", "期初预支余额",
        "本次抵扣", "期末预支余额", "实发金额",
    ]
    ws.append(headers)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8EEF5")
    center = Alignment(horizontal="center")
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = bold
        cell.fill = fill
        cell.alignment = center

    for line in lines:
        ws.append([
            line["participant_name"],
            round(line["net_profit"], 2),
            round(line["share"], 2),
            round(line["advance_before"], 2),
            round(line["deducted"], 2),
            round(line["advance_after"], 2),
            round(line["payout"], 2),
        ])

    if lines:
        total_row = ws.max_row + 1
        ws.cell(row=total_row, column=1, value="合计").font = bold
        for col_idx in range(2, 8):
            letter = ws.cell(row=1, column=col_idx).column_letter
            cell = ws.cell(
                row=total_row, column=col_idx,
                value=f"=SUM({letter}2:{letter}{total_row - 1})",
            )
            cell.font = bold

    widths = [12, 14, 14, 14, 12, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
